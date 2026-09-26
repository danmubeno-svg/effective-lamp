#!/usr/bin/env python3
"""Génère Code.gs (Google Apps Script) à partir du classeur Excel d'état des lieux.

Usage : python3 generate.py [chemin_du_classeur.xlsx]

Le script lit chaque onglet de collecte, en extrait les rubriques et produit
un fichier Code.gs autonome (spécification JSON + logique de construction)
à coller dans https://script.google.com pour créer les formulaires Google.
"""
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
DEFAULT_XLSX = HERE / "Formulaires_Etat_des_lieux_PARSS_PAR-MNM_2026-2027.xlsx"


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def read_instructions(ws):
    """Onglet « Mode d'emploi » : consignes générales + ancrage de chaque formulaire."""
    title = clean(ws["A1"].value)
    instructions = []
    anchors = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        a, b, c = (clean(v) for v in row[:3])
        if isinstance(row[0], (int, float)) and b:
            anchors[b] = c
        elif a and b and a not in ("N°",):
            instructions.append({"label": a, "text": b})
    return title, instructions, anchors


def read_equipment(ws):
    families = []
    for row in ws.iter_rows(min_row=7, values_only=True):
        num, family, label, unit = row[0], clean(row[1]), clean(row[2]), clean(row[3])
        if not isinstance(num, (int, float)) or not label:
            continue
        if not families or families[-1]["name"] != family:
            families.append({"name": family, "items": []})
        service, _, unite = unit.partition("|")
        families[-1]["items"].append(
            {"num": int(num), "label": label, "service": service.strip(), "unit": unite.strip()}
        )
    return families


def read_thematic(ws):
    """Onglets à colonnes : N°, Domaine, Champ/question, Unité/modalité, puis E–I à renseigner."""
    items = []
    for row in ws.iter_rows(min_row=7, values_only=True):
        num, domain, label, unit = row[0], clean(row[1]), clean(row[2]), clean(row[3])
        if not label or label.startswith("Validation"):
            continue
        if unit.strip(". ") == "":
            unit = ""
        if isinstance(num, (int, float)):
            items.append({"num": int(num), "domain": domain, "label": label, "unit": unit, "sub": []})
        elif items:
            # Ligne complémentaire rattachée à la rubrique précédente (ex. Infrastructures n°1).
            items[-1]["sub"].append(label)
    for item in items:
        if item["sub"]:
            # La première ligne est elle-même une sous-question.
            item["sub"].insert(0, item["label"])
            item["label"] = item["domain"]
    return items


def read_indicators(ws):
    subtitle = clean(ws["A2"].value)
    items = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        code = clean(row[0])
        if not code:
            continue
        items.append(
            {
                "code": code,
                "level": clean(row[1]),
                "label": clean(row[2]),
                "unit": clean(row[6]),
                "source": clean(row[8]),
                "comment": clean(row[9]),
            }
        )
    return subtitle, items


def build_spec(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path)
    sheets = wb.worksheets
    title, instructions, anchors = read_instructions(sheets[0])
    forms = []
    for ws in sheets[1:]:
        name = ws.title
        heading = clean(ws["A1"].value)
        if name == "Équipements":
            forms.append({"key": name, "kind": "equipment", "title": heading,
                          "anchor": anchors.get(name, ""), "families": read_equipment(ws)})
        elif name == "Suivi indicateurs":
            subtitle, items = read_indicators(ws)
            forms.append({"key": name, "kind": "indicators", "title": heading,
                          "anchor": subtitle, "items": items})
        else:
            forms.append({"key": name, "kind": "thematic", "title": heading,
                          "anchor": anchors.get(name, ""), "items": read_thematic(ws)})
    return {
        "title": title,
        "period": "PARSS-SSR / PAR-MNM | Période visée : 2026–2027",
        "instructions": instructions,
        "forms": forms,
    }


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    spec = build_spec(xlsx)
    logic = (HERE / "FormBuilder.gs").read_text(encoding="utf-8")
    header = (
        "/**\n"
        " * FICHIER GÉNÉRÉ par generate.py à partir de :\n"
        f" *   {xlsx.name}\n"
        " * Ne pas modifier SPEC à la main : modifier le classeur puis relancer generate.py.\n"
        " */\n\n"
    )
    body = "const SPEC = " + json.dumps(spec, ensure_ascii=False, indent=2) + ";\n\n"
    out = HERE / "Code.gs"
    out.write_text(header + body + logic, encoding="utf-8")
    counts = []
    for f in spec["forms"]:
        n = sum(len(fam["items"]) for fam in f["families"]) if f["kind"] == "equipment" else len(f["items"])
        counts.append(f"  - {f['key']}: {n} rubriques")
    print(f"Écrit {out} ({len(spec['forms'])} formulaires)\n" + "\n".join(counts))


if __name__ == "__main__":
    main()
